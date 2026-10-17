"""Download and read source files: csv/tsv, xlsx, or a zip holding a delimited file.

Two design rules, both worth defending:

1. No pandas here. pandas guesses types: an ABN "05123456789" becomes the
   integer 5123456789 (leading zero gone), a charity named "NA" becomes NaN.
   The engine wants every cell as a string; typing happens later in transforms,
   where it is validated.

2. The reader never guesses the format. The agent's inspect step looks at the
   raw first lines and writes format / delimiter / header_row / sheet into the
   config; this module just obeys the config. A "CSV" that is really
   tab-separated is simply a config whose delimiter is a tab.
"""

import csv
import zipfile
from pathlib import Path

import httpx
import openpyxl

from .config import RAW, USER_AGENT


def download(url: str, source_id: str) -> Path:
    """Stream a URL into data/raw/<source_id>/. Skips if already downloaded."""
    folder = RAW / source_id
    folder.mkdir(parents=True, exist_ok=True)
    name = url.split("/")[-1].split("?")[0]
    if name == "":
        name = "download.bin"
    target = folder / name
    if target.exists() and target.stat().st_size > 0:
        return target
    partial = target.with_suffix(target.suffix + ".part")
    with httpx.stream(
        "GET", url, headers={"User-Agent": USER_AGENT}, timeout=300, follow_redirects=True
    ) as resp:
        resp.raise_for_status()
        with partial.open("wb") as fh:
            for chunk in resp.iter_bytes():
                fh.write(chunk)
    partial.rename(target)  # atomic: an interrupted download is never cached
    return target


def unzip_if_needed(path: Path) -> Path:
    """If path is a .zip, extract its largest csv/tsv/txt member once and return it.

    Extracting to disk costs some space but keeps everything after this point
    a plain text file — the simplest honest way to handle zips.
    """
    if path.suffix.lower() != ".zip":
        return path
    with zipfile.ZipFile(path) as zf:
        candidates = []
        for info in zf.infolist():
            if info.filename.lower().endswith((".csv", ".tsv", ".txt")):
                candidates.append(info)
        if not candidates:
            raise RuntimeError(f"no csv/tsv/txt member inside {path.name}")
        # the largest member is almost always the data file
        biggest = max(candidates, key=lambda info: info.file_size)
        out = path.parent / Path(biggest.filename).name
        if not out.exists():
            zf.extract(biggest, path=path.parent)
            extracted = path.parent / biggest.filename
            if extracted != out:
                extracted.rename(out)  # flatten any folder path inside the zip
        return out


def head_text(path: Path, n_lines: int = 50) -> str:
    """The raw first lines of a text file, for the agent's inspect step."""
    lines = []
    with path.open(encoding="utf-8-sig", errors="replace") as fh:
        for line in fh:
            lines.append(line.rstrip("\n"))
            if len(lines) >= n_lines:
                break
    return "\n".join(lines)


def head_xlsx(path: Path, n_rows: int = 15) -> dict[str, list[list[str]]]:
    """Sheet name → first rows (lists of cell strings), for the inspect step."""
    wb = openpyxl.load_workbook(path, read_only=True)
    sheets = {}
    for name in wb.sheetnames:
        rows = []
        for row in wb[name].iter_rows(values_only=True):
            cells = []
            for cell in row:
                if cell is None:
                    cells.append("")
                else:
                    cells.append(str(cell))
            rows.append(cells)
            if len(rows) >= n_rows:
                break
        sheets[name] = rows
    wb.close()
    return sheets


def clean_header(cells: list) -> list[str]:
    """Strip header names, call blank columns col_N, suffix duplicates with _2, _3…"""
    header = []
    used: set[str] = set()
    for i, cell in enumerate(cells):
        if cell is None:
            name = ""
        else:
            name = str(cell).strip()
        if name == "":
            name = f"col_{i}"
        # keep bumping until the name is genuinely unused — a plain "_2" could
        # collide with a real column already named "A_2", silently dropping it
        base = name
        n = 1
        while name in used:
            n += 1
            name = f"{base}_{n}"
        used.add(name)
        header.append(name)
    return header


def read_rows(path: Path, reader_cfg: dict, limit: int) -> list[dict]:
    """Read up to `limit` records as {column: string} dicts, following the config.

    reader_cfg comes from the agent-written mapping config, e.g.
    {"format": "csv", "encoding": "utf-8", "delimiter": ",", "header_row": 0}
    or {"format": "xlsx", "sheet": "Income tax details", "header_row": 1}.
    """
    records = []
    for record in iter_rows(path, reader_cfg):
        records.append(record)
        if len(records) >= limit:
            break  # stop early: a 250MB file is never fully read
    return records


def iter_rows(path: Path, reader_cfg: dict):
    """Yield records one by one (used to scan big files without loading them)."""
    fmt = reader_cfg["format"]
    if fmt in ("csv", "tsv"):
        yield from _iter_delimited(path, reader_cfg)
    elif fmt == "xlsx":
        yield from _iter_xlsx(path, reader_cfg)
    else:
        raise RuntimeError(f"unknown format in reader config: {fmt}")


def _row_to_record(header: list[str], cells: list) -> dict:
    """Pair header names with cell values; pad short rows, drop extra cells."""
    record = {}
    for i, name in enumerate(header):
        if i < len(cells):
            value = cells[i]
            if value is None:
                value = ""
            record[name] = str(value)
        else:
            record[name] = ""  # row was shorter than the header
    return record


def _iter_delimited(path: Path, cfg: dict):
    encoding = cfg.get("encoding", "utf-8-sig")
    header_row = cfg.get("header_row", 0)
    with path.open(encoding=encoding, errors="replace", newline="") as fh:
        reader = csv.reader(fh, delimiter=cfg["delimiter"])
        for _ in range(header_row):  # skip junk lines above the header
            next(reader)
        header = clean_header(next(reader))
        for cells in reader:
            yield _row_to_record(header, cells)


def _iter_xlsx(path: Path, cfg: dict):
    header_row = cfg.get("header_row", 0)
    wb = openpyxl.load_workbook(path, read_only=True)  # read_only streams rows
    if cfg["sheet"] not in wb.sheetnames:
        wb.close()
        raise RuntimeError(f"sheet {cfg['sheet']!r} not found; sheets: {wb.sheetnames}")
    header: list[str] | None = None
    for i, row in enumerate(wb[cfg["sheet"]].iter_rows(values_only=True)):
        if i < header_row:
            continue  # junk preamble above the header
        if header is None:
            header = clean_header(list(row))
            continue
        yield _row_to_record(header, list(row))
    wb.close()
